#!/usr/bin/env python3
"""Validate Excel structures that openpyxl can write but Excel will repair."""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


def _bounds(reference: str) -> tuple[int, int, int, int]:
    bounds = range_boundaries(reference)
    if any(value is None for value in bounds):
        raise ValueError(f"range must identify concrete cells: {reference}")
    return bounds  # type: ignore[return-value]


def ranges_overlap(left: str, right: str) -> bool:
    """Return whether two same-sheet A1 ranges share at least one cell."""
    left_min_col, left_min_row, left_max_col, left_max_row = _bounds(left)
    right_min_col, right_min_row, right_max_col, right_max_row = _bounds(right)
    return not (
        left_max_col < right_min_col
        or right_max_col < left_min_col
        or left_max_row < right_min_row
        or right_max_row < left_min_row
    )


def validate_workbook(filename: str | Path) -> dict[str, Any]:
    """Check an XLSX/XLSM workbook for Excel-incompatible filter overlaps."""
    path = Path(filename)
    if not path.is_file():
        return {
            "status": "error",
            "file": str(path),
            "message": "file does not exist",
            "issues": [],
        }

    try:
        workbook = load_workbook(
            path,
            data_only=False,
            read_only=False,
            keep_vba=path.suffix.lower() == ".xlsm",
        )
    except Exception as exc:
        return {
            "status": "error",
            "file": str(path),
            "message": f"workbook could not be read: {exc}",
            "issues": [],
        }

    issues: list[dict[str, str]] = []
    try:
        for worksheet in workbook.worksheets:
            worksheet_filter = str(worksheet.auto_filter.ref or "").strip()
            if not worksheet_filter:
                continue

            for table in worksheet.tables.values():
                table_range = str(table.ref or "").strip()
                if not table_range:
                    continue
                try:
                    overlaps = ranges_overlap(worksheet_filter, table_range)
                except (TypeError, ValueError) as exc:
                    issues.append(
                        {
                            "code": "invalid_filter_range",
                            "sheet": worksheet.title,
                            "worksheet_filter": worksheet_filter,
                            "table": table.name,
                            "table_range": table_range,
                            "message": str(exc),
                        }
                    )
                    continue
                if overlaps:
                    issues.append(
                        {
                            "code": "worksheet_table_filter_overlap",
                            "sheet": worksheet.title,
                            "worksheet_filter": worksheet_filter,
                            "table": table.name,
                            "table_range": table_range,
                            "message": (
                                "worksheet AutoFilter overlaps an Excel table; "
                                "remove worksheet.auto_filter.ref or move it outside the table"
                            ),
                        }
                    )
    finally:
        workbook.close()

    return {
        "status": "invalid" if issues else "valid",
        "file": str(path.resolve()),
        "issues": issues,
    }


def main(argv: list[str] | None = None) -> int:
    args = list(sys.argv[1:] if argv is None else argv)
    if len(args) != 1:
        print(
            json.dumps(
                {
                    "status": "error",
                    "message": "usage: validate_xlsx.py <excel_file>",
                    "issues": [],
                }
            )
        )
        return 2

    result = validate_workbook(args[0])
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return {"valid": 0, "invalid": 1}.get(result["status"], 2)


if __name__ == "__main__":
    raise SystemExit(main())
