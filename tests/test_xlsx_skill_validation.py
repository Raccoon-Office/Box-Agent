from __future__ import annotations

import importlib.util
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from box_agent.tools.skill_loader import SkillLoader


REPO_ROOT = Path(__file__).resolve().parents[1]
XLSX_SKILL_DIR = REPO_ROOT / "box_agent" / "skills" / "document-skills" / "xlsx"
VALIDATOR_PATH = XLSX_SKILL_DIR / "scripts" / "validate_xlsx.py"


def _load_validator():
    spec = importlib.util.spec_from_file_location(
        "box_agent_xlsx_validator", VALIDATOR_PATH
    )
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


VALIDATOR = _load_validator()


def _save_workbook(
    path: Path,
    *,
    worksheet_filter: str | None = None,
    table_range: str | None = None,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "People"
    for row in (
        ("Name", "Role", "City", "Department", "Status", "Owner"),
        ("Ada", "Engineer", "London", "R&D", "Active", "A"),
        ("Lin", "Designer", "Shanghai", "Design", "Active", "B"),
        ("Sam", "Analyst", "New York", "Finance", "Active", "C"),
    ):
        worksheet.append(row)
    if worksheet_filter is not None:
        worksheet.auto_filter.ref = worksheet_filter
    if table_range is not None:
        worksheet.add_table(Table(displayName="PeopleTable", ref=table_range))
    workbook.save(path)
    workbook.close()


@pytest.mark.parametrize(
    ("worksheet_filter", "table_range"),
    [
        (None, "A1:C4"),
        ("A1:C4", None),
        ("E1:F4", "A1:C4"),
    ],
)
def test_validator_accepts_non_overlapping_filter_structures(
    tmp_path: Path,
    worksheet_filter: str | None,
    table_range: str | None,
) -> None:
    output = tmp_path / "valid.xlsx"
    _save_workbook(
        output,
        worksheet_filter=worksheet_filter,
        table_range=table_range,
    )

    result = VALIDATOR.validate_workbook(output)

    assert result["status"] == "valid"
    assert result["issues"] == []


@pytest.mark.parametrize("worksheet_filter", ["A1:C4", "B1:D4"])
def test_validator_rejects_worksheet_filter_overlapping_table(
    tmp_path: Path, worksheet_filter: str
) -> None:
    output = tmp_path / "invalid.xlsx"
    _save_workbook(
        output,
        worksheet_filter=worksheet_filter,
        table_range="A1:C4",
    )

    result = VALIDATOR.validate_workbook(output)

    assert result["status"] == "invalid"
    assert result["issues"] == [
        {
            "code": "worksheet_table_filter_overlap",
            "sheet": "People",
            "worksheet_filter": worksheet_filter,
            "table": "PeopleTable",
            "table_range": "A1:C4",
            "message": (
                "worksheet AutoFilter overlaps an Excel table; "
                "remove worksheet.auto_filter.ref or move it outside the table"
            ),
        }
    ]


def test_validator_reports_missing_workbook(tmp_path: Path) -> None:
    missing = tmp_path / "missing.xlsx"

    result = VALIDATOR.validate_workbook(missing)

    assert result == {
        "status": "error",
        "file": str(missing),
        "message": "file does not exist",
        "issues": [],
    }


def test_xlsx_skill_requires_excel_compatibility_validation() -> None:
    instructions = (XLSX_SKILL_DIR / "SKILL.md").read_text(encoding="utf-8")

    assert "Never set `worksheet.auto_filter.ref`" in instructions
    assert "python scripts/validate_xlsx.py <excel_file>" in instructions
    assert "openpyxl.load_workbook()" in instructions


def test_xlsx_skill_resolves_validator_to_an_executable_absolute_path() -> None:
    loader = SkillLoader(REPO_ROOT / "box_agent" / "skills")
    loader.discover_skills()

    skill = loader.get_skill("xlsx")

    assert skill is not None
    assert str(VALIDATOR_PATH) in skill.content
