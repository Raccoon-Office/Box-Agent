"""Subject identity survives the independent design handoff and local repairs."""
import json

from tests.test_pptx_design_plan import design_case, run, write, record_response, scaffold
from tests.test_pptx_design_recovery import record_correction
from tests.test_pptx_design_regressions import js, render, probe


def test_designer_receives_visual_intent_keywords_and_subject_recommendation(design_case, monkeypatch):
    import base64
    monkeypatch.setenv('BOX_AGENT_SOURCE_TEXT_B64', base64.b64encode('李白：生平与诗歌成就'.encode()).decode())
    root, outline, _, _ = design_case
    outline['deck_goal'] = '李白：生平与诗歌成就'
    outline['slides'][0]['visual'] = '水墨山月与人物剪影'
    write(root / 'outline.json', outline)
    prepared = run('design_plan.js', 'prepare', 'outline.json', cwd=root)
    assert prepared.returncode == 0, prepared.stderr
    input_data = json.loads((root / 'design_input.json').read_text())
    brief = json.loads(open(input_data['request_file']).read())
    pages = [row for file in brief['content_files'] for row in json.loads(open(file).read())]
    themes = [row for file in brief['theme_index_files'] for row in json.loads(open(file).read())]
    assert pages[0]['visual_hint'] == '水墨山月与人物剪影'
    assert next(t for t in themes if t['id'] == 'soft-editorial')['mood_keywords']
    assert brief['theme_recommendation']['theme_id'] == 'soft-editorial'
    assert brief['theme_recommendation']['confidence'] == 'high'
    assert brief['subject_profile']['motifs'] == ['ink-landscape']
    outline['slides'][0]['visual'] = '明月与孤帆'
    write(root / 'outline.json', outline)
    assert run('design_plan.js', 'prepare', 'outline.json', cwd=root).returncode == 0
    assert json.loads((root / 'design_input.json').read_text())['input_hash'] == input_data['input_hash']
    assert json.loads(open(input_data['request_file']).read()) == brief


def test_subject_profile_does_not_override_colors_or_user_opt_out(tmp_path):
    data = js(tmp_path, r'''
const path=require('path'),root=process.argv[2];
const plans=require(path.join(root,'scripts/design_plan_core.js'));
const subject=require(path.join(root,'scripts/subject_visual_core.js'));
const outline={deck_goal:'李白与唐诗',slides:[{title:'诗仙',message:'山水与豪情',bullets:['明月','孤帆']}]};
const input=plans.makeInput(outline,'李白','李白与唐诗，白底、蓝色强调');
const decision={theme_id:'soft-editorial',palette:{background:'#FFFFFF',text:'#111111',primary:'#2563EB',accent:'#2563EB',secondary:'#71814B',accent_usage:'sparse'},visual_requirements:{canvas:'solid',heading:'editorial',display_font:'serif',body_font:'sans-serif',shadow:'none',allow_plain_fallback:true},slides:[{layout_id:'cards-grid-v1'}]};
const plan=plans.canonicalPlan(decision,input);
console.log(JSON.stringify({profile:plan.visual_profile,palette:plan.palette,
 optout:subject.inferSubjectProfile({outline,source_text:'李白介绍，不要水墨'}),
 optouts:['李白介绍，不要使用水墨','李白介绍，不使用传统风格',"Li Bai, do not use Chinese ink",'李白介绍，采用漫画风'].map(source_text=>subject.inferSubjectProfile({outline,source_text})),
 rejectedComic:subject.inferSubjectProfile({outline,source_text:'李白介绍，不要漫画'}),
 generic:subject.inferSubjectProfile({title:'季度经营复盘'}),
 disabledCss:subject.profileCss(plan.visual_profile,true)}));
''')
    assert data['profile']['id'] == 'classical-poetry'
    assert data['palette']['primary'] == '#2563EB'
    assert data['palette']['background'] == '#FFFFFF'
    assert data['optout'] is None and data['generic'] is None
    assert all(profile is None for profile in data['optouts'])
    assert data['rejectedComic']['id'] == 'classical-poetry'
    assert data['disabledCss'] == ''


def decision_from(plan):
    decision = json.loads(json.dumps({key: plan[key] for key in ['theme_id', 'palette', 'visual_requirements', 'reason']}))
    decision['slides'] = [{key: slide[key] for key in ['layout_id', 'visual_options']} for slide in json.loads(json.dumps(plan['slides']))]
    return decision


def test_unbound_second_design_cannot_replace_valid_palette_or_layout(design_case):
    root, _, _, plan = design_case
    changed = decision_from(plan)
    changed['palette'] = {**plan['palette'], 'accent': '#A35F82'}
    changed['slides'] = [{**s, 'visual_options': {'composition': 'standard'}} for s in changed['slides']]
    record_response(root, changed)
    result = run('design_plan.js', 'accept', 'design_input.json', cwd=root)
    assert result.returncode == 0, result.stderr
    assert json.loads((root / 'design_plan.json').read_text()) == plan


def test_designer_workspace_alias_does_not_discard_completed_design(design_case):
    root, outline, _, plan = design_case
    alias = root.parent / (root.name + '-alias')
    try:
        alias.symlink_to(root, target_is_directory=True)
    except OSError:
        import pytest
        pytest.skip('Directory symlinks are unavailable')
    outline['deck_goal'] = '李白的课堂介绍'
    write(root / 'outline.json', outline)
    assert run('design_plan.js', 'prepare', 'outline.json', cwd=root).returncode == 0
    sid = record_response(root, decision_from(plan), workspace=alias)
    accepted = run('design_plan.js', 'accept', 'design_input.json', cwd=root)
    assert accepted.returncode == 0, accepted.stderr
    assert json.loads(accepted.stdout)['source_session'] == sid
    assert json.loads((root / 'design_plan.json').read_text())['visual_profile']['id'] == 'classical-poetry'


def test_page_correction_preserves_palette_subject_and_all_other_pages(design_case):
    root, _, _, plan = design_case
    packet = run('design_plan.js', 'correct', 'design_input.json', '--pages', '2', '--issue', 'Page 2 overflows', cwd=root)
    assert packet.returncode == 0, packet.stderr
    from pathlib import Path
    packet_path = Path(json.loads(packet.stdout)['correction_file'])
    changed = decision_from(plan)
    changed['palette'] = {**plan['palette'], 'accent': '#A35F82'}
    changed['theme_id'] = 'soft-editorial'
    changed['slides'] = [{**s, 'visual_options': {'composition': 'standard'}} for s in changed['slides']]
    record_correction(root, packet_path, changed)
    result = run('design_plan.js', 'accept', 'design_input.json', cwd=root)
    assert result.returncode == 0, result.stderr
    accepted = json.loads((root / 'design_plan.json').read_text())
    assert accepted['palette'] == plan['palette']
    assert accepted['theme_id'] == plan['theme_id']
    assert accepted['slides'][0] == plan['slides'][0]
    assert accepted['slides'][2] == plan['slides'][2]
    assert accepted['slides'][1]['visual_options'] == {'composition': 'standard'}
    assert scaffold(root).returncode == 0


def test_poetry_profile_is_rendered_on_every_slide_without_extra_media(design_case):
    root, outline, _, plan = design_case
    outline['deck_goal'] = '李白的生平与诗歌成就'
    write(root / 'outline.json', outline)
    assert run('design_plan.js', 'prepare', 'outline.json', cwd=root).returncode == 0
    d = decision_from(plan)
    record_response(root, d)
    assert run('design_plan.js', 'accept', 'design_input.json', cwd=root).returncode == 0
    assert scaffold(root).returncode == 0
    deck_path = root / 'deck.json'
    deck = json.loads(deck_path.read_text())
    for slide, page in zip(deck['slides'], outline['slides']):
        slide['props']['title'] = page['title']
        slide['props']['subtitle'] = page['message']
        slide['props']['items'] = [{'kicker':str(i+1),'title':b,'body':b} for i,b in enumerate(page['bullets'])]
    html = render(root, deck_path, deck)
    text = html.read_text()
    assert 'data-deck-profile="classical-poetry"' in text
    assert 'data-deck-profile-motifs="ink-landscape"' in text
    assert 'data:image/svg+xml' in text
    assert all('data:image/svg+xml' in row['motif'] and row['motifDisplay'] != 'none' for row in browser_styles(root, html))
    qa = probe(html)
    assert qa['editor']['componentContrast']['failureCount'] == 0


def test_featured_five_items_rejected_before_rendering_and_correction_keeps_other_pages(design_case):
    root, outline, _, plan = design_case
    outline['tone'] = 'Five item regression'
    outline['slides'][1]['bullets'] = ['One','Two','Three','Four','Five']
    write(root / 'outline.json', outline)
    assert run('design_plan.js', 'prepare', 'outline.json', cwd=root).returncode == 0
    original = decision_from(plan)
    original['slides'][1]['visual_options']['variant'] = 'featured'
    record_response(root, original)
    failed = run('design_plan.js', 'accept', 'design_input.json', cwd=root)
    assert failed.returncode != 0
    assert 'featured supports at most 4' in failed.stderr
    from pathlib import Path
    packet = Path(json.loads(failed.stderr)['correction_file'])
    changed = decision_from(plan)
    changed['slides'] = [{**s,'visual_options':{'composition':'standard','variant':'balanced'}} for s in changed['slides']]
    changed['palette'] = {**plan['palette'],'accent':'#A35F82'}
    record_correction(root, packet, changed)
    accepted = run('design_plan.js','accept','design_input.json',cwd=root)
    assert accepted.returncode == 0, accepted.stderr
    new_plan = json.loads((root/'design_plan.json').read_text())
    assert new_plan['slides'][0] == plan['slides'][0]
    assert new_plan['slides'][2] == plan['slides'][2]
    assert new_plan['palette'] == plan['palette']


def browser_styles(tmp_path, html):
    return js(tmp_path, r'''
const path=require('path'),{pathToFileURL}=require('url');
const host=require(path.join(process.argv[2],'scripts/playwright_host.js'));
host.ensurePlaywrightBrowsersPath();
const {chromium}=host.loadPlaywright();
(async()=>{const browser=await chromium.launch(host.chromiumLaunchOptions(chromium,{headless:true}).options);
try {const page=await browser.newPage();await page.goto(pathToFileURL(process.argv[3]).href);
await page.evaluate(()=>window.__deckTextReady);
console.log(JSON.stringify(await page.evaluate(()=>Array.from(document.querySelectorAll('#deck-root > .slide')).map(s=>({
 motif:getComputedStyle(s,'::after').backgroundImage,
 motifDisplay:getComputedStyle(s,'::after').display,
 wash:s.querySelector('.slide-background')?getComputedStyle(s.querySelector('.slide-background'),'::after').opacity:null,
 color:s.querySelector('.slide-background')?getComputedStyle(s.querySelector('.slide-background'),'::after').backgroundColor:null
})))));
}finally{await browser.close();}})().catch(e=>{console.error(e);process.exit(1)});
''', html)


def test_frozen_palette_preserves_visible_image_and_dark_treatment(tmp_path):
    from tests.test_pptx_design_regressions import render_case
    from tests.test_pptx_palette_contract import evaluate
    contract = evaluate(tmp_path, "console.log(JSON.stringify(decide('背景色 #FFFFFF，正文色 #1D1D1F，主色 #1D1D1F').contract));")
    deck_path, deck = render_case(tmp_path, ['cover-editorial-v1','cover-editorial-v1'], 'plain-neutral')
    deck['design_contract'] = contract
    image = tmp_path / 'landscape.svg'
    image.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="1920" height="1080"><rect width="1920" height="1080" fill="#71814B"/></svg>')
    for slide, treatment in zip(deck['slides'], ['wash-light','wash-dark']):
        slide['background'] = {'src':image.name,'alt':'Landscape','treatment':treatment,'fit':'cover'}
    html = render(tmp_path, deck_path, deck)
    styles = browser_styles(tmp_path, html)
    assert float(styles[0]['wash']) == 0.72
    assert float(styles[1]['wash']) == 0.82
    import re
    assert sum(map(int, re.findall(r'\d+', styles[1]['color']))) < 96


def test_timeline_numbers_retain_room_for_two_digits_in_editable_export(tmp_path):
    from tests.test_pptx_design_regressions import render_case
    from pptx import Presentation
    deck_path, deck = render_case(tmp_path, ['timeline-horizontal-v1'], 'soft-editorial')
    html = render(tmp_path, deck_path, deck)
    exported = run('html_to_editable_pptx.js', html, tmp_path / 'deck.pptx', '--out', tmp_path / 'previews', cwd=tmp_path)
    assert exported.returncode == 0, exported.stdout + exported.stderr
    indices = [shape for shape in Presentation(tmp_path / 'deck.pptx').slides[0].shapes
               if shape.has_text_frame and shape.text == '01'
               and shape.text_frame.paragraphs[0].runs[0].font.size.pt > 20]
    assert indices
    for shape in indices:
        # Exact browser glyph bounds leave no margin for a viewer's fallback
        # metrics. Two digits need explicit width in the native text box.
        font_size = shape.text_frame.paragraphs[0].runs[0].font.size.pt
        assert shape.width / 914400 * 72 >= font_size * 1.2


def test_subject_identity_guides_asset_prompts_and_survives_designer_refinement(design_case, monkeypatch):
    import base64
    monkeypatch.setenv('BOX_AGENT_SOURCE_TEXT_B64', base64.b64encode('制作李白与唐诗的介绍，简洁清晰'.encode()).decode())
    root, outline, _, plan = design_case
    outline['deck_goal'] = '唐诗与李白'
    outline['slides'][0].update(title='李白与唐诗',layout='cover',visual='水墨山月与人物剪影')
    write(root/'outline.json',outline)
    assert run('design_plan.js','prepare','outline.json',cwd=root).returncode == 0
    d = decision_from(plan)
    d['slides'][0] = {'layout_id':'cover-editorial-v1','visual_options':{'composition':'poster'}}
    d['visual_profile'] = {'id':'poetic-moon','motifs':['moon'],'semantic_tags':['李白']}
    record_response(root,d)
    accepted = run('design_plan.js','accept','design_input.json',cwd=root)
    assert accepted.returncode == 0, accepted.stderr
    assert scaffold(root).returncode == 0
    deck = json.loads((root/'deck.json').read_text())
    assert deck['design_plan']['visual_profile']['motifs'] == ['ink-landscape','moon']
    manifest = json.loads((root/'assets/generated/manifest.json').read_text())
    prompt = manifest['image_plan'][0]['prompt']
    assert '水墨山月与人物剪影' in prompt
    assert 'Subject visual identity: Conceptual Chinese ink landscape' in prompt
